#!/usr/bin/env python3
"""Build a Korean 16:9 deck from slides with title, bullets, notes and shapes."""
import argparse
import json
from pathlib import Path
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from platform_paths import find_korean_font

DRAFT = '※ 초안입니다. 수업 전 반드시 검토하세요.'
FONT = 'Malgun Gothic'


def apply_run_font(run, size):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string('172B4D')
    props = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        element = OxmlElement(tag)
        element.set('typeface', FONT)
        props.append(element)


def textbox(slide, x, y, width, height, lines, size):
    frame = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    ).text_frame
    frame.word_wrap = True
    for index, text in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.space_after = Pt(10)
        for run in paragraph.runs:
            apply_run_font(run, size)


def set_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            apply_run_font(run, 16)
            run.font.color.rgb = RGBColor.from_string('172B4D')


def validate(spec):
    if not isinstance(spec, dict) or not isinstance(spec.get('slides'), list) or not spec['slides']:
        raise ValueError('slides must be a nonempty list')
    for slide in spec['slides']:
        if not isinstance(slide, dict):
            raise ValueError('each slide must be an object')
        if not isinstance(slide.get('title'), str) or not slide['title'].strip():
            raise ValueError('title must be a nonempty string')
        if not isinstance(slide.get('notes'), str):
            raise ValueError('notes must be a string')
        if not isinstance(slide.get('bullets'), list) or not all(isinstance(b, str) for b in slide['bullets']):
            raise ValueError('bullets must be a list of strings')
        shapes = slide.get('shapes', [])
        if not isinstance(shapes, list):
            raise ValueError('shapes must be a list')
        if len(shapes) > 3:
            raise ValueError('Use at most three fraction bars per slide; split across slides.')
        for shape in shapes:
            if not isinstance(shape, dict):
                raise ValueError('each shape must be an object')
            cells, filled = shape.get('cells'), shape.get('filled')
            if type(cells) is not int or type(filled) is not int or not 0 <= filled <= cells or cells < 1:
                raise ValueError('fraction bars require integer cells > 0 and 0 <= filled <= cells')
            color = shape.get('color', '176B87')
            if not isinstance(color, str) or not re.fullmatch(r'[0-9a-fA-F]{6}', color):
                raise ValueError('color must be six RGB hex digits')


def main():
    global FONT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    if args.output.exists():
        parser.error(f'Output already exists: {args.output}')
    font = find_korean_font()
    FONT = font.family
    spec = json.loads(args.spec.read_text(encoding='utf-8'))
    validate(spec)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    for data in spec['slides']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        textbox(slide, 0.65, 0.35, 12, 1.1, [data['title']], 34)
        bars = data.get('shapes', [])
        textbox(slide, 0.8, 1.6, 11.7, 2.4 if bars else 4.8, data['bullets'], 24)
        for row, bar in enumerate(bars):
            width = 10.8 / bar['cells']
            for cell in range(bar['cells']):
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, Inches(1.1 + width * cell),
                    Inches(4.25 + row * 0.65), Inches(width), Inches(0.48)
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(
                    bar.get('color', '176B87') if cell < bar['filled'] else 'FFFFFF'
                )
                shape.line.color.rgb = RGBColor.from_string('172B4D')
        set_notes(slide, data['notes'])
    textbox(deck.slides[-1], 0.65, 6.85, 12, 0.5, [DRAFT], 20)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.output)
    print(f'pptx={args.output.resolve()} slides={len(deck.slides)} font={FONT} font_file={font.path or "NOT_FOUND"}')


if __name__ == '__main__':
    main()
